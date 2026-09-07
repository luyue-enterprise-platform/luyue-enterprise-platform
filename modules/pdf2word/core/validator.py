# -*- coding: utf-8 -*-
"""转换校验器（v2.1.0 新增）

转换完成后对输出文件逐项校验，覆盖用户六条验收要求：
1. orientation  页面方向：输出与源文件逐段方向一致
2. page_size    页面尺寸：输出所有页面/节均为 A4
3. layout_text  布局结构-文本：输出文本对源文本覆盖率 ≥ 阈值（防截断/错位丢字）
4. layout_images 布局结构-图片：图片数量比对
5. layout_tables 布局结构-表格：表格数量比对
6. fonts        字体格式：字体/字号/粗斜体/颜色签名覆盖率 ≥ 阈值

校验不通过时的两级处理：
- 可自动修正项（docx 节方向/尺寸、PDF 页面规范化）→ 自动修正后重新校验（status=fixed）
- 仍不通过 → status=fail，detail 注明具体文件名、页码/段落、不一致项

报告结构（每文件一份）：
{
  'file': 源文件名, 'out': 输出文件名, 'direction': 'pdf2word'|'topdf',
  'overall': 'pass'|'fixed'|'fail',
  'items': [{'key', 'label', 'status': 'pass'|'fixed'|'fail', 'detail'}...]
}
"""
import os
import re
import zipfile
from collections import Counter

import fitz  # PyMuPDF

from . import content_orient, page_norm

# ---------------- 阈值常量 ----------------
TEXT_COVERAGE_MIN = 0.95      # 文本覆盖率（防截断丢字）
FONT_COVERAGE_MIN = 0.85      # 字体签名覆盖率（按字符数加权）
IMAGE_RATIO_MIN = 0.85        # docx 图片数 / PDF 图片块数
TABLE_RATIO_MIN = 0.80        # docx 表格数 / PDF 检测表格数
FONT_SIZE_TOL = 1.5           # 字号容差（pt）
COLOR_CHANNEL_TOL = 24        # 颜色每通道容差
EMU_PER_PT = 12700            # docx EMU 换算


# ==================== 文本提取与覆盖率 ====================

_WS_RE = re.compile(r'\s+')


def _norm_text(t):
    return _WS_RE.sub('', t or '')


def _coverage(ref_text, out_text):
    """out 对 ref 的字符多重集覆盖率：sum(min(c_ref,c_out)) / sum(c_ref)"""
    ref = Counter(_norm_text(ref_text))
    if not ref:
        return 1.0
    out = Counter(_norm_text(out_text))
    hit = sum(min(cnt, out.get(ch, 0)) for ch, cnt in ref.items())
    return hit / max(1, sum(ref.values()))


def pdf_text(pdf_path):
    parts = []
    with fitz.open(pdf_path) as doc:
        for page in doc:
            parts.append(page.get_text())
    return '\n'.join(parts)


def docx_text(docx_path):
    from docx import Document
    doc = Document(docx_path)
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            if hf is not None:
                parts.extend(p.text for p in hf.paragraphs)
    return '\n'.join(parts)


# ==================== 字体签名比对 ====================

def _clean_font_name(name):
    """去 PDF 子集前缀（ABCDEF+SimSun）+ 归一化"""
    if not name:
        return ''
    name = re.sub(r'^[A-Z]{6}\+', '', name)
    return name.replace(' ', '').lower()


def _font_flags(span_flags):
    """PyMuPDF span flags: bit1=italic(2), bit4=bold(16)"""
    return bool(span_flags & 16), bool(span_flags & 2)


def pdf_font_signatures(pdf_path):
    """PDF 文本 span 签名 -> 字符数权重: {(font,size,bold,italic,color): chars}"""
    sigs = Counter()
    with fitz.open(pdf_path) as doc:
        for page in doc:
            for block in page.get_text('dict').get('blocks', []):
                if block.get('type') != 0:
                    continue
                for line in block.get('lines', []):
                    for span in line.get('spans', []):
                        text = _norm_text(span.get('text', ''))
                        if not text:
                            continue
                        bold, italic = _font_flags(span.get('flags', 0))
                        key = (_clean_font_name(span.get('font', '')),
                               round(float(span.get('size', 0)), 1),
                               bold, italic,
                               int(span.get('color', 0)))
                        sigs[key] += len(text)
    return sigs


def docx_font_signatures(docx_path):
    """docx run 签名 -> 字符数权重"""
    from docx import Document
    doc = Document(docx_path)
    sigs = Counter()

    def _eat(paragraphs):
        for p in paragraphs:
            for run in p.runs:
                text = _norm_text(run.text)
                if not text:
                    continue
                fname = _clean_font_name(run.font.name or '')
                size = round(run.font.size.pt, 1) if run.font.size else 0.0
                bold = bool(run.font.bold)
                italic = bool(run.font.italic)
                color = 0
                try:
                    if run.font.color is not None and run.font.color.rgb is not None:
                        color = int(str(run.font.color.rgb), 16)
                except Exception:
                    color = -1
                sigs[(fname, size, bold, italic, color)] += len(text)

    _eat(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                _eat(cell.paragraphs)
    return sigs


def _sig_match(pdf_key, docx_keys):
    """PDF 签名在 docx 签名集合中找匹配（字号/颜色容差，粗斜体严格）"""
    pf, ps, pb, pi, pc = pdf_key
    for df, ds, db, di, dc in docx_keys:
        if pb != db or pi != di:
            continue
        if pf and df and pf != df:
            continue
        if ps and ds and abs(ps - ds) > FONT_SIZE_TOL:
            continue
        if pc >= 0 and dc >= 0:
            dr = abs(((pc >> 16) & 255) - ((dc >> 16) & 255))
            dg = abs(((pc >> 8) & 255) - ((dc >> 8) & 255))
            db2 = abs((pc & 255) - (dc & 255))
            if max(dr, dg, db2) > COLOR_CHANNEL_TOL:
                continue
        return True
    return False


def font_coverage(ref_sigs, out_sigs):
    """ref 签名按字符加权，在 out 中的命中比例"""
    total = sum(ref_sigs.values())
    if total == 0:
        return 1.0
    out_keys = list(out_sigs.keys())
    hit = sum(cnt for key, cnt in ref_sigs.items() if _sig_match(key, out_keys))
    return hit / total


# ==================== 图片/表格计数 ====================

def pdf_image_count(pdf_path):
    n = 0
    with fitz.open(pdf_path) as doc:
        for page in doc:
            try:
                n += len(page.get_image_info())
            except Exception:
                n += len(page.get_images())
    return n


def docx_image_count(docx_path):
    with zipfile.ZipFile(docx_path) as zf:
        return len([n for n in zf.namelist()
                    if n.startswith('word/media/') and not n.endswith('/')])


def pdf_table_count(pdf_path):
    n = 0
    with fitz.open(pdf_path) as doc:
        for page in doc:
            try:
                n += len(page.find_tables().tables)
            except Exception:
                pass
    return n


def docx_table_count(docx_path):
    from docx import Document
    return len(Document(docx_path).tables)


# ==================== docx 节方向/尺寸 ====================

def docx_sections_info(docx_path):
    """[(orientation, width_pt, height_pt), ...] 按文档顺序

    方向以页面实际宽高为准（w>h 即横版）——w:orient 属性不可靠：
    pdf2docx 只写 pgSz 宽高、不写 orient 属性，Word 渲染同样以宽高为准。
    """
    from docx import Document
    out = []
    for sec in Document(docx_path).sections:
        w = (sec.page_width or 0) / EMU_PER_PT
        h = (sec.page_height or 0) / EMU_PER_PT
        orient = page_norm.page_orientation(w, h)
        out.append((orient, w, h))
    return out


def fix_docx_sections_to_a4(docx_path, expected_orientations):
    """自动修正：把 docx 各节页面设为 A4 并按期望方向段对齐方向。

    expected_orientations: 源文件逐页方向（run-length 与 docx 节一一对应；
    节数多于段数时按最后一段方向补齐，段数多于节数时截断）。
    同时写 pgSz 宽高与 orient 属性（Word 以宽高为准，属性保持自洽）。
    修正后保存。返回修正的节数。
    """
    from docx import Document
    from docx.enum.section import WD_ORIENT
    doc = Document(docx_path)
    runs = page_norm.orientation_runs(expected_orientations)
    sections = doc.sections
    fixed = 0
    for i, sec in enumerate(sections):
        orient = runs[i][0] if i < len(runs) else runs[-1][0]
        w_pt, h_pt = page_norm.a4_size_for(orient)
        w_emu, h_emu = int(w_pt * EMU_PER_PT), int(h_pt * EMU_PER_PT)
        if abs((sec.page_width or 0) - w_emu) > EMU_PER_PT or \
           abs((sec.page_height or 0) - h_emu) > EMU_PER_PT:
            fixed += 1
        sec.page_width = w_emu
        sec.page_height = h_emu
        sec.orientation = (WD_ORIENT.LANDSCAPE if orient == 'landscape'
                           else WD_ORIENT.PORTRAIT)
    if sections:
        doc.save(docx_path)
    return fixed


# ==================== 报告条目工具 ====================

def _item(key, label, status, detail):
    return {'key': key, 'label': label, 'status': status, 'detail': detail}


def _overall(items):
    statuses = {i['status'] for i in items}
    if 'fail' in statuses:
        return 'fail'
    if 'fixed' in statuses:
        return 'fixed'
    return 'pass'


def make_report(file_name, out_name, direction, items):
    return {
        'file': file_name,
        'out': out_name,
        'direction': direction,
        'overall': _overall(items),
        'items': items,
    }


# ==================== 方向/尺寸校验（PDF 输出） ====================

def _check_pdf_pages(pdf_path, expected_orientations, file_label):
    """校验 PDF 输出：逐页 A4 + 逐页方向。返回 (items, bad_pages)"""
    size_bad, orient_bad = [], []
    with fitz.open(pdf_path) as doc:
        for i, page in enumerate(doc):
            w, h = page.rect.width, page.rect.height
            if not page_norm.is_a4(w, h):
                size_bad.append(i + 1)
            if i < len(expected_orientations):
                actual = page_norm.page_orientation(w, h)
                if actual != expected_orientations[i]:
                    orient_bad.append((i + 1, expected_orientations[i], actual))
    items = []
    if orient_bad:
        pages = '、'.join('第%d页(应%s/实%s)' % (p, '横版' if e == 'landscape' else '竖版',
                          '横版' if a == 'landscape' else '竖版')
                          for p, e, a in orient_bad[:5])
        items.append(_item('orientation', '页面方向', 'fail',
                           '%s 方向不一致：%s%s' % (file_label, pages,
                           ' 等' if len(orient_bad) > 5 else '')))
    else:
        items.append(_item('orientation', '页面方向', 'pass',
                           '共 %d 页方向与源文件一致' % len(expected_orientations)))
    if size_bad:
        items.append(_item('page_size', '页面尺寸(A4)', 'fail',
                           '%s 第 %s 页非 A4 尺寸' % (
                               file_label, '、'.join(str(p) for p in size_bad[:5]))))
    else:
        items.append(_item('page_size', '页面尺寸(A4)', 'pass', '全部页面为 A4'))
    return items, (size_bad, orient_bad)


def normalize_pdf_until_valid(pdf_path, expected_orientations, file_label):
    """PDF 输出自动修正：尺寸/方向不过 → 再规范化一次重验。返回 items"""
    items, (size_bad, orient_bad) = _check_pdf_pages(pdf_path, expected_orientations, file_label)
    if not size_bad and not orient_bad:
        return items
    # 自动修正：整体再走一遍 A4 规范化（方向以源文件期望为准逐页强制）
    tmp = pdf_path + '.renorm.pdf'
    try:
        out_doc = fitz.open()
        with fitz.open(pdf_path) as src_doc:
            for i, page in enumerate(src_doc):
                orient = (expected_orientations[i] if i < len(expected_orientations)
                          else page_norm.page_orientation(page.rect.width, page.rect.height))
                dst_w, dst_h = page_norm.a4_size_for(orient)
                new_page = out_doc.new_page(width=dst_w, height=dst_h)
                r = page.rect
                scale = min(dst_w / r.width, dst_h / r.height)
                show_w, show_h = r.width * scale, r.height * scale
                x0, y0 = (dst_w - show_w) / 2.0, (dst_h - show_h) / 2.0
                new_page.show_pdf_page(fitz.Rect(x0, y0, x0 + show_w, y0 + show_h),
                                       src_doc, page.number)
        out_doc.save(tmp, garbage=3, deflate=True)
        out_doc.close()
        os.replace(tmp, pdf_path)
    except Exception:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        return items  # 修正失败，保留原 fail 项
    # 重新校验
    re_items, (sb2, ob2) = _check_pdf_pages(pdf_path, expected_orientations, file_label)
    fixed_items = []
    for old, new in zip(items, re_items):
        if old['status'] == 'fail' and new['status'] == 'pass':
            fixed_items.append(_item(new['key'], new['label'], 'fixed',
                                     '已自动修正并复核通过（%s）' % new['detail']))
        else:
            fixed_items.append(new)
    return fixed_items


# ==================== pdf2word 方向校验 ====================

def validate_pdf_to_word(norm_pdf_path, docx_path, src_orientations, file_name, out_name):
    """PDF→Word 转换校验（基准：A4 规范化后的 PDF）。

    流程：方向/尺寸校验 → 不一致自动修正 docx 节 → 重新校验 → 其余项校验。
    """
    items = []

    # ---- 1. 页面方向 + 2. 页面尺寸（docx 节） ----
    expected_runs = page_norm.orientation_runs(src_orientations)
    sections = docx_sections_info(docx_path)
    orient_mismatch, size_bad = [], []
    for i, (orient, w, h) in enumerate(sections):
        exp = expected_runs[i][0] if i < len(expected_runs) else expected_runs[-1][0]
        if orient != exp:
            seg = expected_runs[i] if i < len(expected_runs) else expected_runs[-1]
            orient_mismatch.append((i + 1, seg, exp, orient))
        if not page_norm.is_a4(w, h):
            size_bad.append(i + 1)

    if orient_mismatch or size_bad:
        # 自动修正：重写全部节为 A4 + 期望方向，保存后重验
        fix_docx_sections_to_a4(docx_path, src_orientations)
        sections2 = docx_sections_info(docx_path)
        om2, sb2 = [], []
        for i, (orient, w, h) in enumerate(sections2):
            exp = expected_runs[i][0] if i < len(expected_runs) else expected_runs[-1][0]
            if orient != exp:
                om2.append(i + 1)
            if not page_norm.is_a4(w, h):
                sb2.append(i + 1)
        if om2:
            seg = expected_runs[om2[0] - 1] if om2[0] - 1 < len(expected_runs) else expected_runs[-1]
            items.append(_item('orientation', '页面方向', 'fail',
                               '%s 第 %d-%d 页方向不一致（自动修正未通过）'
                               % (file_name, seg[1], seg[2])))
        elif orient_mismatch:
            items.append(_item('orientation', '页面方向', 'fixed',
                               '检测到方向/尺寸偏差已自动修正，复核通过（%d 段）'
                               % len(expected_runs)))
        else:
            items.append(_item('orientation', '页面方向', 'pass',
                               '%d 个方向段与源文件一致' % len(expected_runs)))
        if sb2:
            items.append(_item('page_size', '页面尺寸(A4)', 'fail',
                               '%s 节 %s 非 A4（自动修正未通过）'
                               % (file_name, '、'.join(map(str, sb2[:5])))))
        elif size_bad:
            items.append(_item('page_size', '页面尺寸(A4)', 'fixed',
                               '已统一修正为 A4 并复核通过'))
        else:
            items.append(_item('page_size', '页面尺寸(A4)', 'pass', '全部页面为 A4'))
    else:
        items.append(_item('orientation', '页面方向', 'pass',
                           '%d 个方向段与源文件一致' % len(expected_runs)))
        items.append(_item('page_size', '页面尺寸(A4)', 'pass', '全部页面为 A4'))

    # ---- 3a. 布局结构：文本完整性（防截断/错位丢字） ----
    ref_text = pdf_text(norm_pdf_path)
    out_text = docx_text(docx_path)
    cov = _coverage(ref_text, out_text)
    items.append(_item(
        'layout_text', '布局结构-文本', 'pass' if cov >= TEXT_COVERAGE_MIN else 'fail',
        '%s 文本覆盖率 %.1f%%（阈值 %.0f%%）'
        % (file_name, cov * 100, TEXT_COVERAGE_MIN * 100)))

    # ---- 3b. 布局结构：图片 ----
    pdf_imgs = pdf_image_count(norm_pdf_path)
    docx_imgs = docx_image_count(docx_path)
    img_ok = (pdf_imgs == 0) or (docx_imgs >= pdf_imgs * IMAGE_RATIO_MIN)
    items.append(_item(
        'layout_images', '布局结构-图片', 'pass' if img_ok else 'fail',
        '%s 图片 %d 张（源 %d 张）' % (file_name, docx_imgs, pdf_imgs)))

    # ---- 3c. 布局结构：表格 ----
    pdf_tbls = pdf_table_count(norm_pdf_path)
    docx_tbls = docx_table_count(docx_path)
    tbl_ok = (pdf_tbls == 0) or (docx_tbls >= pdf_tbls * TABLE_RATIO_MIN)
    items.append(_item(
        'layout_tables', '布局结构-表格', 'pass' if tbl_ok else 'fail',
        '%s 表格 %d 个（源检出 %d 个）' % (file_name, docx_tbls, pdf_tbls)))

    # ---- 4. 字体与格式 ----
    pdf_sigs = pdf_font_signatures(norm_pdf_path)
    docx_sigs = docx_font_signatures(docx_path)
    fcov = font_coverage(pdf_sigs, docx_sigs)
    items.append(_item(
        'fonts', '字体与格式', 'pass' if fcov >= FONT_COVERAGE_MIN else 'fail',
        '%s 字体格式签名覆盖率 %.1f%%（阈值 %.0f%%）'
        % (file_name, fcov * 100, FONT_COVERAGE_MIN * 100)))

    return make_report(file_name, out_name, 'pdf2word', items)


# ==================== topdf 方向校验 ====================

def validate_word_to_pdf(docx_path, pdf_path, file_name, out_name):
    """Word→PDF：方向段（docx 节 vs PDF 页段）、A4、文本、字体"""
    sections = docx_sections_info(docx_path)
    sec_orients = [s[0] for s in sections]
    # PDF 逐页方向段与 docx 节方向段比对（run-length 序列前缀应一致）
    with fitz.open(pdf_path) as doc:
        page_orients = [page_norm.page_orientation(p.rect.width, p.rect.height) for p in doc]
    page_runs = page_norm.orientation_runs(page_orients)
    expected = []
    for orient, _, _ in [(r[0], r[1], r[2]) for r in page_runs]:
        expected.append(orient)
    sec_seq = sec_orients[:len(expected)] if len(sec_orients) >= len(expected) else sec_orients
    # 页段可能因重排比节多（空节被吞），只校验每页方向 ∈ 对应节方向：退化为整体一致性
    mismatch = len(page_runs) > 0 and sec_seq and any(
        pr[0] != sec_seq[min(i, len(sec_seq) - 1)] for i, pr in enumerate(page_runs))

    items, _ = _check_pdf_pages(pdf_path, page_orients, file_name)  # 先按自身方向验尺寸
    # 用节方向期望重验方向项
    items = [i for i in items if i['key'] != 'orientation']
    if mismatch:
        items.insert(0, _item('orientation', '页面方向', 'fail',
                              '%s PDF 页方向段与 Word 节方向不一致' % file_name))
    else:
        items.insert(0, _item('orientation', '页面方向', 'pass',
                              '页方向段与 Word 节方向一致（%d 段）' % len(page_runs)))

    cov = _coverage(docx_text(docx_path), pdf_text(pdf_path))
    items.append(_item('layout_text', '布局结构-文本',
                       'pass' if cov >= TEXT_COVERAGE_MIN else 'fail',
                       '%s 文本覆盖率 %.1f%%（阈值 %.0f%%）'
                       % (file_name, cov * 100, TEXT_COVERAGE_MIN * 100)))
    fcov = font_coverage(docx_font_signatures(docx_path), pdf_font_signatures(pdf_path))
    items.append(_item('fonts', '字体与格式',
                       'pass' if fcov >= FONT_COVERAGE_MIN else 'fail',
                       '%s 字体格式签名覆盖率 %.1f%%（阈值 %.0f%%）'
                       % (file_name, fcov * 100, FONT_COVERAGE_MIN * 100)))
    return make_report(file_name, out_name, 'topdf', items)


def validate_excel_to_pdf(xlsx_path, pdf_path, file_name, out_name):
    """Excel→PDF：A4 尺寸 + 文本覆盖（单元格值 vs PDF 文本）"""
    items, _ = _check_pdf_pages_self(pdf_path, file_name)
    ref = ''
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        vals = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                vals.extend(str(v) for v in row if v is not None)
        ref = ' '.join(vals)
        wb.close()
    except Exception:
        ref = ''
    if ref:
        cov = _coverage(ref, pdf_text(pdf_path))
        items.append(_item('layout_text', '布局结构-文本',
                           'pass' if cov >= TEXT_COVERAGE_MIN else 'fail',
                           '%s 文本覆盖率 %.1f%%（阈值 %.0f%%）'
                           % (file_name, cov * 100, TEXT_COVERAGE_MIN * 100)))
    return make_report(file_name, out_name, 'topdf', items)


def _check_pdf_pages_self(pdf_path, file_label):
    """只验 A4 尺寸与方向自洽（方向以生成时设定为准，这里回读作为 expected）"""
    with fitz.open(pdf_path) as doc:
        orients = [page_norm.page_orientation(p.rect.width, p.rect.height) for p in doc]
    return _check_pdf_pages(pdf_path, orients, file_label)


def validate_image_to_pdf(image_paths, pdf_path, file_name, out_name,
                          expected_orients=None):
    """图片→PDF：页数=图片数、A4、方向符合主体内容预期、每页含图。

    expected_orients：转换阶段已判定的逐图方向（v2.1.2，避免重复 OCR）；缺省时
    按主体内容（EXIF 归一化 + OCR 文字区域，对右下角水印鲁棒）自行判定。
    """
    from PIL import Image
    expected = list(expected_orients) if expected_orients else []
    if not expected:
        for p in image_paths:
            with Image.open(p) as im:
                if im.mode in ('RGBA', 'P', 'LA'):
                    im = im.convert('RGB')
                _n, o, _b = content_orient.decide_orientation(im)
            expected.append(o)
    items, _ = _check_pdf_pages(pdf_path, expected, file_name)
    with fitz.open(pdf_path) as doc:
        cnt = doc.page_count
        no_img_pages = [i + 1 for i, page in enumerate(doc)
                        if not page.get_image_info()]
    items.append(_item(
        'layout_images', '布局结构-图片',
        'pass' if (cnt == len(image_paths) and not no_img_pages) else 'fail',
        '%s 输出 %d 页/源 %d 张%s'
        % (file_name, cnt, len(image_paths),
           '，第 %s 页无图片' % '、'.join(map(str, no_img_pages[:5])) if no_img_pages else '')))
    return make_report(file_name, out_name, 'topdf', items)


def validate_text_to_pdf(txt_path, pdf_path, file_name, out_name):
    """TXT→PDF：A4 + 文本覆盖"""
    items, _ = _check_pdf_pages_self(pdf_path, file_name)
    raw = ''
    for enc in ('utf-8', 'gbk', 'utf-16'):
        try:
            with open(txt_path, 'r', encoding=enc) as f:
                raw = f.read()
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    cov = _coverage(raw, pdf_text(pdf_path))
    items.append(_item('layout_text', '布局结构-文本',
                       'pass' if cov >= TEXT_COVERAGE_MIN else 'fail',
                       '%s 文本覆盖率 %.1f%%（阈值 %.0f%%）'
                       % (file_name, cov * 100, TEXT_COVERAGE_MIN * 100)))
    return make_report(file_name, out_name, 'topdf', items)
