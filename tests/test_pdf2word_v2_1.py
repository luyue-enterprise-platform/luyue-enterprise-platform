# -*- coding: utf-8 -*-
"""pdf2word v2.1.0 双向保真 + 逐项校验测试

六条验收要求映射：
1. 页面方向：normalize 逐页方向保持；validate orientation 项
2. 页面尺寸：全部 A4 + 内容等比缩放不变形不截断；validate page_size 项
3. 布局结构：文本/图片/表格校验项（layout_text/layout_images/layout_tables）
4. 字体格式：字体签名覆盖率校验项（fonts）
5. 转换校验：每文件转换后逐项校验，失败自动修正（fixed）或明确报错（fail 含文件名页码）
6. 批量处理：results 逐文件携带 validation 报告，前端逐项展示
"""
import os
import sys
import shutil
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import fitz
from modules.pdf2word.core import page_norm, validator, to_pdf
from modules.pdf2word.core.converter import convert_pdf_to_docx


def _make_mixed_pdf(path, with_image=True):
    """构造测试 PDF：第1页竖版A5文本（含粗体/颜色/不同字号），第2页横版文本+图片"""
    doc = fitz.open()
    # 第 1 页：竖版小页（A5 竖 420×595）
    p1 = doc.new_page(width=420, height=595)
    p1.insert_text((50, 80), '第一页竖版标题', fontsize=18, fontname='china-s',
                   color=(0.8, 0, 0))
    p1.insert_text((50, 120), '正文内容 abcdefg 1234567 保真校验专用文本', fontsize=11,
                   fontname='china-s', color=(0, 0, 0))
    p1.insert_text((50, 160), '第二段：方向保持与 A4 规范化验证', fontsize=11,
                   fontname='china-s')
    # 第 2 页：横版（595×420）
    p2 = doc.new_page(width=595, height=420)
    p2.insert_text((60, 90), '第二页横版内容 landscape page', fontsize=14,
                   fontname='china-s', color=(0, 0, 0.8))
    p2.insert_text((60, 130), '横版页面转换后必须保持横版', fontsize=11,
                   fontname='china-s')
    if with_image:
        # 生成一张小图插入第 2 页
        from PIL import Image
        import io
        img = Image.new('RGB', (80, 50), (30, 120, 200))
        buf = io.BytesIO()
        img.save(buf, 'PNG')
        p2.insert_image(fitz.Rect(60, 160, 140, 210), stream=buf.getvalue())
    doc.save(path)
    doc.close()
    return path


class TestPageNorm(unittest.TestCase):
    """要求1+2：方向逐页保持 + 统一 A4 + 等比缩放不变形"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_orientation_helpers(self):
        self.assertEqual(page_norm.page_orientation(595, 842), 'portrait')
        self.assertEqual(page_norm.page_orientation(842, 595), 'landscape')
        self.assertEqual(page_norm.page_orientation(500, 500), 'portrait')  # 方形按竖
        self.assertTrue(page_norm.is_a4(*page_norm.A4_PORTRAIT))
        self.assertTrue(page_norm.is_a4(*page_norm.A4_LANDSCAPE))
        self.assertFalse(page_norm.is_a4(612, 792))  # Letter 非 A4

    def test_orientation_runs(self):
        runs = page_norm.orientation_runs(
            ['portrait', 'portrait', 'landscape', 'portrait'])
        self.assertEqual(runs, [('portrait', 1, 2), ('landscape', 3, 3),
                                ('portrait', 4, 4)])

    def test_normalize_mixed_pdf(self):
        src = _make_mixed_pdf(os.path.join(self.tmp, 'mixed.pdf'))
        dst = os.path.join(self.tmp, 'norm.pdf')
        info = page_norm.normalize_pdf_to_a4(src, dst)
        self.assertEqual(info['pages'], 2)
        self.assertEqual(info['orientations'], ['portrait', 'landscape'])
        with fitz.open(dst) as doc:
            # 第 1 页 A4 竖
            r1 = doc[0].rect
            self.assertTrue(page_norm.is_a4(r1.width, r1.height))
            self.assertLess(r1.width, r1.height)
            # 第 2 页 A4 横
            r2 = doc[1].rect
            self.assertTrue(page_norm.is_a4(r2.width, r2.height))
            self.assertGreater(r2.width, r2.height)
            # 文本未栅格化（仍可提取）
            self.assertIn('第一页竖版标题', doc[0].get_text())
            self.assertIn('横版页面转换后必须保持横版', doc[1].get_text())
            # 内容等比缩放：scale = min(宽比, 高比)，不变形
            s1 = info['scales'][0]
            self.assertAlmostEqual(s1, min(page_norm.A4_W / 420, page_norm.A4_H / 595), places=4)


class TestValidatorHelpers(unittest.TestCase):

    def test_coverage(self):
        self.assertAlmostEqual(validator._coverage('abcdef', 'abcdef'), 1.0)
        self.assertAlmostEqual(validator._coverage('aabbcc', 'aabb'), 4 / 6, places=3)
        self.assertEqual(validator._coverage('', 'anything'), 1.0)
        # 空白归一化
        self.assertEqual(validator._coverage('a b\nc', 'abc'), 1.0)

    def test_clean_font_name(self):
        self.assertEqual(validator._clean_font_name('ABCDEF+SimSun'), 'simsun')
        self.assertEqual(validator._clean_font_name('Microsoft YaHei'), 'microsoftyahei')
        self.assertEqual(validator._clean_font_name(''), '')

    def test_sig_match_tolerance(self):
        # 字号容差 ±1.5，颜色通道容差 24，粗斜体严格
        pdf_key = ('simsun', 11.0, True, False, 0x102030)
        self.assertTrue(validator._sig_match(pdf_key, [('simsun', 12.0, True, False, 0x102032)]))
        self.assertFalse(validator._sig_match(pdf_key, [('simsun', 14.0, True, False, 0x102030)]))
        self.assertFalse(validator._sig_match(pdf_key, [('simsun', 11.0, False, False, 0x102030)]))
        self.assertFalse(validator._sig_match(pdf_key, [('simsun', 11.0, True, True, 0x102030)]))


class TestFixDocxSections(unittest.TestCase):
    """要求5：自动修正——docx 节尺寸/方向改 A4 并可复核"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_fix_letter_to_a4_landscape(self):
        from docx import Document
        path = os.path.join(self.tmp, 't.docx')
        d = Document()  # 默认 Letter 竖版
        d.add_paragraph('x')
        d.save(path)
        fixed = validator.fix_docx_sections_to_a4(path, ['landscape'])
        self.assertGreaterEqual(fixed, 1)
        infos = validator.docx_sections_info(path)
        self.assertEqual(infos[0][0], 'landscape')
        self.assertTrue(page_norm.is_a4(infos[0][1], infos[0][2]))


class TestPdf2WordValidation(unittest.TestCase):
    """端到端：PDF→规范化→docx→逐项校验（要求1-5全链路）"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_convert_with_validation_report(self):
        src = _make_mixed_pdf(os.path.join(self.tmp, '样例.pdf'))
        dst = os.path.join(self.tmp, '样例.docx')
        r = convert_pdf_to_docx(src, dst)
        self.assertTrue(r['ok'], r.get('error'))
        self.assertEqual(r['pages'], 2)
        rep = r['validation']
        self.assertIsNotNone(rep)
        self.assertEqual(rep['direction'], 'pdf2word')
        self.assertIn(rep['overall'], ('pass', 'fixed'), rep)
        # 逐项存在且含关键校验项
        keys = [i['key'] for i in rep['items']]
        for k in ('orientation', 'page_size', 'layout_text',
                  'layout_images', 'layout_tables', 'fonts'):
            self.assertIn(k, keys)
        # 方向/尺寸必须过（pass 或 fixed）
        for i in rep['items']:
            if i['key'] in ('orientation', 'page_size'):
                self.assertIn(i['status'], ('pass', 'fixed'),
                              '%s: %s' % (i['key'], i['detail']))
        # 文本覆盖率项必须通过（防截断丢字）
        lt = [i for i in rep['items'] if i['key'] == 'layout_text'][0]
        self.assertEqual(lt['status'], 'pass', lt['detail'])
        # docx 节均为 A4 且方向段与源一致
        infos = validator.docx_sections_info(dst)
        self.assertEqual([s[0] for s in infos], ['portrait', 'landscape'])
        for orient, w, h in infos:
            self.assertTrue(page_norm.is_a4(w, h))

    def test_report_fail_detail_contains_filename(self):
        """要求5：不一致时 detail 注明文件名（构造文本大量缺失的坏 docx 触发 fail）"""
        src = _make_mixed_pdf(os.path.join(self.tmp, '断字.pdf'))
        norm = os.path.join(self.tmp, 'norm.pdf')
        info = page_norm.normalize_pdf_to_a4(src, norm)
        # 造一个几乎空白的 docx
        from docx import Document
        bad = os.path.join(self.tmp, '断字.docx')
        d = Document()
        d.add_paragraph('仅存留极少文本')
        d.save(bad)
        rep = validator.validate_pdf_to_word(norm, bad, info['orientations'],
                                             '断字.pdf', '断字.docx')
        self.assertEqual(rep['overall'], 'fail')
        lt = [i for i in rep['items'] if i['key'] == 'layout_text'][0]
        self.assertEqual(lt['status'], 'fail')
        self.assertIn('断字.pdf', lt['detail'])  # 明确注明文件名


class TestToPdfValidation(unittest.TestCase):
    """逆向转换同规则：图片/文本/PDF 透传的 A4 + 校验"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_image_landscape_to_a4_landscape(self):
        from PIL import Image
        src = os.path.join(self.tmp, '横图.png')
        Image.new('RGB', (1600, 900), (200, 30, 30)).save(src)
        dst = os.path.join(self.tmp, '横图.pdf')
        pages, orient = to_pdf.convert_image_to_pdf(src, dst)
        self.assertEqual(pages, 1)
        self.assertEqual(orient, 'landscape')
        with fitz.open(dst) as doc:
            r = doc[0].rect
            self.assertTrue(page_norm.is_a4(r.width, r.height))
            self.assertGreater(r.width, r.height)  # 横图→横版
        rep = validator.validate_image_to_pdf([src], dst, '横图.png', '横图.pdf')
        self.assertEqual(rep['overall'], 'pass', rep['items'])

    def test_image_portrait_to_a4_portrait(self):
        from PIL import Image
        src = os.path.join(self.tmp, '竖图.png')
        Image.new('RGB', (600, 1200), (30, 200, 30)).save(src)
        dst = os.path.join(self.tmp, '竖图.pdf')
        to_pdf.convert_image_to_pdf(src, dst)
        with fitz.open(dst) as doc:
            r = doc[0].rect
            self.assertLess(r.width, r.height)  # 竖图→竖版
        rep = validator.validate_image_to_pdf([src], dst, '竖图.png', '竖图.pdf')
        self.assertEqual(rep['overall'], 'pass', rep['items'])

    def test_text_to_pdf_validation(self):
        src = os.path.join(self.tmp, '说明.txt')
        with open(src, 'w', encoding='utf-8') as f:
            f.write('逆向转换保真校验文本\n第二行内容 abc 123\n' * 20)
        dst = os.path.join(self.tmp, '说明.pdf')
        pages, rep = to_pdf.convert_text_to_pdf(src, dst)
        self.assertGreaterEqual(pages, 1)
        self.assertEqual(rep['overall'], 'pass', rep['items'])
        keys = [i['key'] for i in rep['items']]
        self.assertIn('page_size', keys)
        self.assertIn('layout_text', keys)

    def test_pdf_passthrough_normalized(self):
        """PDF 透传也走 A4 规范化：Letter 尺寸输入 → 全 A4 输出 + 报告通过"""
        src = os.path.join(self.tmp, 'letter.pdf')
        doc = fitz.open()
        p = doc.new_page(width=612, height=792)  # US Letter 竖版
        p.insert_text((72, 100), 'Letter size content 透传规范化', fontsize=12,
                      fontname='china-s')
        doc.save(src)
        doc.close()
        dst = os.path.join(self.tmp, 'letter_out.pdf')
        pages, rep = to_pdf.convert_pdf_to_pdf(src, dst)
        self.assertEqual(pages, 1)
        with fitz.open(dst) as d2:
            r = d2[0].rect
            self.assertTrue(page_norm.is_a4(r.width, r.height))
            self.assertLess(r.width, r.height)  # 方向保持竖版
        self.assertEqual(rep['overall'], 'pass', rep['items'])

    def test_batch_results_carry_validation(self):
        """要求6：批量独立模式 results 逐文件携带 validation 报告"""
        from PIL import Image
        f1 = os.path.join(self.tmp, 'a.png')
        Image.new('RGB', (800, 500), (10, 10, 10)).save(f1)
        f2 = os.path.join(self.tmp, 'b.txt')
        with open(f2, 'w', encoding='utf-8') as f:
            f.write('批量校验测试文本内容')
        out = os.path.join(self.tmp, 'out')
        results, skipped = to_pdf.batch_to_pdf([f1, f2], out,
                                               output_mode='individual')
        self.assertEqual(len(results), 2)
        self.assertEqual(skipped, [])
        for r in results:
            self.assertIsNotNone(r['validation'])
            self.assertIn(r['validation']['overall'], ('pass', 'fixed'),
                          r['validation']['items'])


class TestBlueprintValidationFlow(unittest.TestCase):
    """端点级：上传 → 后台转换 → 结果接口逐项携带 validation（要求5+6 链路）"""

    def setUp(self):
        from app import app as flask_app
        self.c = flask_app.test_client()
        with self.c.session_transaction() as sess:
            sess['user_id'] = 1
            sess['username'] = 'tester'

    def _png(self):
        import io as _io
        from PIL import Image
        buf = _io.BytesIO()
        Image.new('RGB', (640, 400), (50, 50, 120)).save(buf, 'PNG')
        return buf.getvalue()

    def test_topdf_result_carries_validation(self):
        import io as _io
        import time
        r = self.c.post('/pdf2word/api/upload', data={
            'direction': 'topdf', 'output_mode': 'individual',
            'files': [(_io.BytesIO(self._png()), '横图.png')],
        }, content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200, r.get_json())
        tid = r.get_json()['task_id']
        for _ in range(100):
            p = self.c.get('/pdf2word/api/progress/%s' % tid).get_json()
            if p['status'] in ('done', 'error'):
                break
            time.sleep(0.2)
        self.assertEqual(p['status'], 'done')
        res = self.c.get('/pdf2word/api/result/%s' % tid).get_json()
        self.assertEqual(len(res['results']), 1)
        v = res['results'][0]['validation']
        self.assertIsNotNone(v)
        self.assertEqual(v['direction'], 'topdf')
        self.assertIn(v['overall'], ('pass', 'fixed'), v['items'])
        keys = [i['key'] for i in v['items']]
        self.assertIn('orientation', keys)
        self.assertIn('page_size', keys)


class TestFrontendValidationUI(unittest.TestCase):
    """前端：校验列 + 明细展开 + 报告导出接线"""

    def test_html_has_validation_column_and_export_btn(self):
        base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        html = open(os.path.join(base, 'modules/pdf2word/templates/pdf2word_index.html'),
                    encoding='utf-8').read()
        self.assertIn('<th>校验</th>', html)
        self.assertIn('exportValidationReport', html)

    def test_js_has_validation_rendering(self):
        base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        js = open(os.path.join(base, 'modules/pdf2word/static/js/app.js'),
                  encoding='utf-8').read()
        for s in ('toggleValDetail', 'exportValidationReport', 'val-badge',
                  'item.validation', '校验通过'):
            self.assertIn(s, js)

    def test_css_has_validation_styles(self):
        base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        css = open(os.path.join(base, 'modules/pdf2word/static/css/style.css'),
                   encoding='utf-8').read()
        for s in ('.val-badge', '.val-pass', '.val-fixed', '.val-fail',
                  '.val-detail-box', '.val-item-label'):
            self.assertIn(s, css)


if __name__ == '__main__':
    unittest.main(verbosity=2)
